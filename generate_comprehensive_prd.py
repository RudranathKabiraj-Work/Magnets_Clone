import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def create_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D1D5DB", sz="4")

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_background(hdr_cells[i], "0066B2") # Brand Primary Blue
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
            run.font.name = "Arial"

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=120, bottom=120, left=160, right=160)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(31, 41, 55)

    # Set Column Widths if provided
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph() # Spacing
    return table

def add_callout_box(doc, title, text, box_type="info"):
    colors = {
        "info": {"bg": "EFF6FF", "border": "3B82F6", "title": "1D4ED8"},
        "note": {"bg": "F3F4F6", "border": "6B7280", "title": "374151"},
        "success": {"bg": "ECFDF5", "border": "10B981", "title": "047857"},
        "warning": {"bg": "FFFBEB", "border": "F59E0B", "title": "B45309"}
    }
    c = colors.get(box_type, colors["info"])

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, c["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{c["border"]}"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(
        int(c["title"][0:2], 16), int(c["title"][2:4], 16), int(c["title"][4:6], 16)
    )

    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(55, 65, 81)

    doc.add_paragraph()

def build_complete_prd():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(31, 41, 55) # Gray-800
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # Document Header / Cover Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("PRODUCT REQUIREMENT DOCUMENT (PRD)\n")
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0, 102, 178) # Brand Accent Blue

    run_title = title_p.add_run("LeadMagnets Platform")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(17, 24, 39)

    tagline_p = doc.add_paragraph()
    tagline_p.paragraph_format.space_after = Pt(18)
    run_tagline = tagline_p.add_run("Next-Generation Lead Magnet Creation, Locked PDF Gating, Social Growth Automation & Lead Nurturing Engine")
    run_tagline.font.size = Pt(11)
    run_tagline.font.italic = True
    run_tagline.font.color.rgb = RGBColor(107, 114, 128)

    # Metadata Table
    create_styled_table(
        doc,
        ["Attribute", "Specification Details"],
        [
            ["Product Name", "LeadMagnets (Magnets Clone Platform)"],
            ["Version / Build", "v1.4.0-Production Ready"],
            ["Document Status", "Approved & Comprehensive Reference Specification"],
            ["Technology Stack", "Next.js 14 (App Router), React 18, TypeScript, Tailwind CSS, MongoDB, Upstash Redis, Three.js, Tiptap"],
            ["Integrations & AI", "Google Gemini 1.5/2.0 AI, Cloudinary, Vercel Blob, Resend, Nodemailer, Unipile/LinkedIn, Calendly, Slack, Pipedrive, Kit"],
            ["Target Ecosystem", "SaaS Founders, B2B Growth Marketers, Content Creators, Lead Gen Agencies, Solopreneurs"],
            ["Date of Documentation", "September 2026"]
        ],
        col_widths=[2.2, 4.3]
    )

    doc.add_page_break()

    # Table of Contents Summary
    h1 = doc.add_heading(level=1)
    r = h1.add_run("Table of Contents")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    toc_items = [
        ("1. Executive Summary & Product Vision", "High-level overview, strategic objectives, value proposition, and competitive advantage."),
        ("2. Target Audience & User Personas", "Detailed personas, pain points, core jobs to be done, and user journey flowcharts."),
        ("3. End-to-End System Architecture", "Frontend, backend, database schemas, edge middleware, AI pipeline, and third-party infrastructure."),
        ("4. Comprehensive Functional Specifications", "Detailed breakdown of all 10 core modules with UI states, workflows, and edge conditions."),
        ("5. Database Schema & Data Models", "Mongoose schemas for Accounts, MagnetPages, Leads, Sequences, Events, OTPs, and Resources."),
        ("6. API Contracts & Webhook Architecture", "REST API endpoints, Cron schedules, and Webhook integrations (LinkedIn, Resend, Calendly)."),
        ("7. UI/UX Design System & Interactive Visuals", "Design philosophy, Dark/Light modes, Three.js 3D canvas, Tiptap editor, and micro-interactions."),
        ("8. Non-Functional Requirements (NFRs)", "Performance benchmarks, security protocols, GDPR/CAN-SPAM compliance, and reliability."),
        ("9. Quality Assurance & Acceptance Criteria", "Comprehensive test scenarios across page builder, OTP gating, sequence cron, and LinkedIn sync."),
        ("10. Product Roadmap & Future Horizons", "Phased implementation milestones, advanced AI roadmap, and multi-tenant scaling.")
    ]

    for title, desc in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run_t = p.add_run(f"• {title}: ")
        run_t.bold = True
        run_t.font.color.rgb = RGBColor(31, 41, 55)
        run_d = p.add_run(desc)
        run_d.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_paragraph()

    # =========================================================================
    # SECTION 1: Executive Summary & Product Vision
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Executive Summary & Product Vision")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    doc.add_paragraph(
        "LeadMagnets is an all-in-one, enterprise-grade lead generation, asset gating, and audience conversion platform. "
        "Historically, creators and growth teams have had to stitch together five disparate tools to execute a single lead magnet campaign: "
        "a landing page builder (e.g., Carrd/Webflow), a file hosting bucket (e.g., Google Drive/Dropbox), an email marketing service (e.g., Mailchimp/ConvertKit), "
        "a social automation scraper (e.g., Phantombuster/Taplio), and an analytics dashboard. This disjointed stack causes friction, dropped leads, "
        "technical maintenance headaches, and high recurring subscription costs."
    )

    doc.add_paragraph(
        "LeadMagnets solves this by providing a unified, high-performance web platform that covers the entire lifecycle of lead acquisition: "
        "from creating sleek, high-converting landing pages in seconds with Google Gemini AI, to offering revolutionary locked-PDF preview gates with email OTP verification, "
        "to automating social comment-to-DM conversions on LinkedIn, to executing intelligent multi-step drip email nurture sequences with stop-on-booking triggers."
    )

    add_callout_box(
        doc,
        "Core Value Proposition",
        "Empower creators and B2B marketers to launch high-converting lead magnet funnels in under 60 seconds, "
        "turn social engagement directly into verified email subscribers, and automatically nurture leads into booked sales calls.",
        "success"
    )

    create_styled_table(
        doc,
        ["Strategic Goal", "Target Metric / KPI", "Platform Mechanism"],
        [
            ["Maximize Conversion Rate", "25% to 45% Visitor-to-Lead conversion", "Minimalist visual templates, A/B testing, and frictionless 1-click forms"],
            ["Zero Deliverable Leakage", "100% verified email captures for PDFs", "Locked PDF in-browser previewer with 6-digit expiring email OTP gate"],
            ["Automate Social Funnels", "5x higher comment-to-lead conversion", "Real-time LinkedIn comment scraping & automated DM delivery engine"],
            ["Drive Sales Meeting Bookings", "40%+ response rate on follow-ups", "Multi-step drip email sequences that automatically stop when Calendly booking occurs"],
            ["Zero Technical Burden", "< 2 minute setup time for custom domains", "Automated DNS verification with seamless Edge Middleware URL rewriting"]
        ],
        col_widths=[2.0, 2.2, 2.3]
    )

    # =========================================================================
    # SECTION 2: Target Audience & User Personas
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("2. Target Audience & User Personas")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    doc.add_paragraph(
        "The platform is specifically engineered for five high-value user archetypes across the digital marketing and creator economy:"
    )

    personas = [
        ["B2B SaaS Founders & Growth Leads", "Needs qualified prospect emails and booked discovery demos.", "Wants fast lead capture pages for whitepapers, calculators, and API docs with Calendly sync."],
        ["LinkedIn Creators & Thought Leaders", "Posts high-engagement content with 'Comment RESOURCE to get the PDF'.", "Tired of manually DMing hundreds of commenters; needs automated comment detection and link delivery."],
        ["Digital Marketing & Lead Gen Agencies", "Manages multiple client campaigns with custom branding.", "Needs rapid page generation, A/B testing, custom domain white-labeling, and webhook CRM forwarding."],
        ["Solopreneurs, Coaches & Info-Product Sellers", "Offers checklists, mini-courses, Notion templates, and video guides.", "Needs zero-code setup, built-in drip email sequences, and aesthetic dark/light page templates."],
        ["Enterprise Sales Teams", "Distributes confidential case studies, teardowns, and proprietary research.", "Requires locked PDF reading gates requiring verified business email OTPs before unlocking full reports."]
    ]

    create_styled_table(doc, ["Persona Archetype", "Core Pain Point / Need", "LeadMagnets Solution"], personas, col_widths=[2.1, 2.2, 2.2])

    # =========================================================================
    # SECTION 3: End-to-End System Architecture
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("3. End-to-End System Architecture")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    doc.add_paragraph(
        "The LeadMagnets platform is built on modern Jamstack & serverless principles using the Next.js 14 App Router, "
        "leveraging Edge Middleware for ultra-low latency routing and authentication validation, and backed by a resilient MongoDB Atlas data layer."
    )

    add_callout_box(
        doc,
        "Architectural Highlights",
        "• Edge-First Authentication: Cryptographic HMAC-SHA256 session token verification in Next.js Edge Runtime using Web Crypto API (crypto.subtle) with 0ms cold starts.\n"
        "• Rate-Limiting Shield: Upstash Redis Token Bucket algorithm protecting sensitive auth and AI endpoints.\n"
        "• High-Performance Rendering: Server-Side Rendering (SSR) for SEO-optimized public magnet pages paired with dynamic client hydration for live preview builders.\n"
        "• Asynchronous Pipeline: Automated Cron runners for processing drip email delays and polling social platforms without blocking web requests.",
        "info"
    )

    tech_stack_data = [
        ["Layer / Subsystem", "Primary Technologies", "Purpose & Architectural Role"],
        ["Frontend Framework", "Next.js 14, React 18, TypeScript", "App router structure, server components, client interactivity, robust type safety."],
        ["Styling & UI Aesthetics", "Tailwind CSS, Framer Motion, Lenis", "Modern glassmorphism, responsive mobile-first layouts, smooth inertial scrolling."],
        ["3D & Canvas Graphics", "Three.js, OGL (Aurora WebGL)", "Interactive 3D magnet canvas on landing hero and dynamic aurora gradient shaders."],
        ["Rich Text Content", "Tiptap Core & Extensions", "WYSIWYG editing for bullet points, pitches, headings, color highlights, and email bodies."],
        ["Document Processing", "pdfjs-dist (v6.3), Docx", "High-fidelity canvas rendering for locked PDF gating and document export generation."],
        ["Backend & API", "Next.js Route Handlers, Node.js", "Serverless REST endpoints, webhook ingestion, and AI orchestrations."],
        ["Database & Caching", "MongoDB Atlas, Mongoose 9.x, Upstash Redis", "Persistent document storage with schema indexing, TTL expiration, and rate limiting."],
        ["AI Copilot", "Google Generative AI (Gemini SDK)", "Zero-shot lead magnet generation, copy optimization, and deliverable personalization."],
        ["Email Delivery", "Resend API, Nodemailer SMTP", "High-deliverability transactional emails, sequence drips, and real-time webhook tracking."],
        ["Storage & Media", "Cloudinary SDK, Vercel Blob", "CDN-backed secure asset storage for lead deliverables, logos, and hero imagery."],
        ["Social Automation", "Unipile API & Custom Webhooks", "LinkedIn account syncing, comment polling, and automated DM message dispatching."]
    ]

    create_styled_table(doc, tech_stack_data[0], tech_stack_data[1:], col_widths=[1.8, 2.3, 2.4])

    # =========================================================================
    # SECTION 4: Comprehensive Functional Specifications
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("4. Comprehensive Functional Specifications")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    doc.add_paragraph(
        "This section details the functional architecture, business logic, workflows, and edge conditions for each of the 10 core modules in the LeadMagnets platform."
    )

    # 4.1 Module 1: Authentication & User Management
    h2 = doc.add_heading(level=2)
    h2.add_run("4.1 Module 1: Authentication, Session Management & Account Setup")
    doc.add_paragraph(
        "• Dual Auth Engine: Supports both NextAuth.js (for provider flexibility) and custom HMAC-SHA256 signed session cookies. "
        "The Edge Middleware parses and verifies tokens using Web Crypto `crypto.subtle` in microseconds without spinning up Node.js runtimes.\n"
        "• Registration & Login: Secure registration with Bcrypt password hashing (salt rounds 10), email uniqueness validation, and instant workspace initialization.\n"
        "• Password Reset Flow: Generates crypto-random reset tokens stored with expiration timestamps in MongoDB, dispatched via Resend transactional email.\n"
        "• Rate-Limiting Defense: Upstash Redis enforces 30 requests per 60 seconds per IP address on `/api/auth/*` and login routes to thwart brute-force attacks."
    )

    # 4.2 Module 2: Visual Lead Magnet Page Builder
    h2 = doc.add_heading(level=2)
    h2.add_run("4.2 Module 2: Visual Lead Magnet Page Builder & Dynamic Templates")
    doc.add_paragraph(
        "The page builder allows users to create high-converting lead magnet landing pages with a real-time reactive split-screen preview:\n"
        "• Template System: Multiple pre-designed layout engines including Minimalist (clean typography), Split-Hero (media side-by-side), Floating Card (elevated glassmorphism), and Dark Modern.\n"
        "• WYSIWYG Content Editing: Powered by Tiptap, enabling inline formatting, custom bullet lists with custom icons, rich pitches, and accented highlight spans.\n"
        "• Asset & Deliverable Management: Users can attach hosted resources (PDFs, Notion templates, checklists, ZIPs) or direct download URLs stored in Cloudinary/Vercel Blob.\n"
        "• Form & Field Customization: Configurable form titles, placeholders, button CTA copy, custom qualifying questions, and multi-field inputs.\n"
        "• After-Signup Funnel Actions: Supports 3 post-conversion destinations:\n"
        "   1. Standard: In-page animated Thank-You state with instant download/access button.\n"
        "   2. Elsewhere: Immediate 302 redirect to a custom URL (e.g., booking page, community link, course portal).\n"
        "   3. Custom: In-page rich media modal featuring embedded Loom/YouTube video, personalized message, and secondary call-to-action button.\n"
        "• A/B Split Testing Engine: Users can enable Variant B to test alternate headlines, subheadlines, and hero imagery. Incoming traffic is evenly split (50/50), with independent impression and conversion counters."
    )

    # 4.3 Module 3: Locked PDF Interactive Viewer
    h2 = doc.add_heading(level=2)
    h2.add_run("4.3 Module 3: Locked PDF Viewer & Email OTP Gating Engine")
    doc.add_paragraph(
        "A marquee feature that provides a superior alternative to standard email gating:\n"
        "• Interactive In-Browser Canvas: Uses `pdfjs-dist` to render crisp vector PDF pages inside the browser without downloading the source file.\n"
        "• Configurable Free Preview: Creators specify how many initial pages (e.g. 1, 2, or 3) are publicly readable.\n"
        "• Blur & Lock Gate: Subsequent pages are covered with an elegant backdrop-blur overlay displaying an embedded email capture form.\n"
        "• Passwordless OTP Verification: When a reader submits their email, the system generates a secure 6-digit numeric OTP code (stored in MongoDB `PdfOtp` collection with a 15-minute TTL index). "
        "The code is dispatched instantly via email. Once verified, the document unlocks entirely in-session, and the user is recorded as a verified lead."
    )

    # 4.4 Module 4: Automated Drip Sequences
    h2 = doc.add_heading(level=2)
    h2.add_run("4.4 Module 4: Automated Drip Email Sequences & Booking Triggers")
    doc.add_paragraph(
        "• Sequence Drip Engine: Multi-step email nurture sequence linked directly to specific lead magnets.\n"
        "• Timed Delays: Granular delay triggers (e.g., Immediately upon signup, +1 hour, +1 day, +3 days, +7 days).\n"
        "• Dynamic Merge Tags: Personalizes email copy with `{{name}}`, `{{magnet_title}}`, `{{deliverable_url}}`, and custom prompt responses.\n"
        "• Stop-on-Booking Automation: When a lead books a sales call via Calendly, a webhook triggers `/api/webhooks/booking`, instantly updating the lead's status to `stopped` and halting all subsequent nurture emails.\n"
        "• Cron Processor: `/api/cron/process-sequences` evaluates pending sequence steps against scheduled delivery times and dispatches emails via Resend/Nodemailer."
    )

    # 4.5 Module 5: LinkedIn Comment-to-Lead Automation
    h2 = doc.add_heading(level=2)
    h2.add_run("4.5 Module 5: LinkedIn Growth & Comment-to-DM Lead Automation")
    doc.add_paragraph(
        "Transforms viral LinkedIn engagement into owned email list subscribers:\n"
        "• Account Connection: Seamless connection via Unipile API / LinkedIn OAuth credentials.\n"
        "• Post Campaign Configuration: Users link specific LinkedIn post IDs/URLs to a designated lead magnet and define trigger keywords (e.g., 'growth', 'guide', 'send').\n"
        "• Automated Comment Ingestion: Polling cron (`/api/cron/linkedin-sync`) and real-time webhooks detect matching comments on creator posts.\n"
        "• Direct Message Dispatch: Automatically sends a personalized LinkedIn DM containing the lead magnet landing link to the commenter.\n"
        "• Anti-Spam Safety: Built-in jitter delays and hourly limits to protect creator LinkedIn accounts from rate-limiting penalties."
    )

    # 4.6 Module 6: Leads Management & CRM
    h2 = doc.add_heading(level=2)
    h2.add_run("4.6 Module 6: Centralized Leads Management & CRM Pipeline")
    doc.add_paragraph(
        "• Complete Lead Dossier: Displays lead name, email, capture source (standard page, locked PDF OTP, LinkedIn DM, integration), device type, referrer URL, and capture timestamp.\n"
        "• Status Lifecycle: Tracks lead journey through `new` → `delivered` → `opened` → `replied` → `converted` / `stopped`.\n"
        "• Tagging & Segmentation: Apply custom tags for categorization and filtering.\n"
        "• Export Engine: 1-click export to formatted CSV and JSON for external reporting or importing into third-party CRMs."
    )

    # 4.7 Module 7: Third-Party Integrations
    h2 = doc.add_heading(level=2)
    h2.add_run("4.7 Module 7: Third-Party Integrations Ecosystem")
    doc.add_paragraph(
        "• Email Marketing: Native sync with ConvertKit / Kit (API v3 subscriber push with tag mapping) and Resend API.\n"
        "• CRMs: Pipedrive CRM API sync for creating deals and contacts upon lead capture.\n"
        "• Real-Time Alerts: Slack Incoming Webhooks (posts instant rich notification cards on every new lead) and Zapier Webhooks.\n"
        "• Scheduling: Calendly integration for meeting tracking and sequence termination.\n"
        "• Analytics: Google Analytics 4 (GA4 Measurement ID) and Meta Pixel ID script injection."
    )

    # 4.8 Module 8: Google Gemini AI Copilot
    h2 = doc.add_heading(level=2)
    h2.add_run("4.8 Module 8: Google Gemini AI Copilot & Deliverable Personalization")
    doc.add_paragraph(
        "• AI Lead Magnet Generator: Users input a topic or target audience; Google Gemini generates complete copy sets including high-converting headlines, subheadings, benefit bullet points, pitch paragraphs, and thank-you delivery emails.\n"
        "• AI Copy Optimizer: One-click enhancement to improve readability, punchiness, and psychological urgency.\n"
        "• Dynamic Deliverable Personalization: Leads answer a custom prompt during signup (e.g., 'What is your current monthly revenue?'); Gemini generates a customized deliverable text block tailored specifically to their input."
    )

    # 4.9 Module 9: Custom Domains & White-Labeling
    h2 = doc.add_heading(level=2)
    h2.add_run("4.9 Module 9: Custom Domains, DNS Verification & White-Labeling")
    doc.add_paragraph(
        "• Custom Subdomains & Apex Domains: Users can map `resources.mybrand.com` or `get.domain.com`.\n"
        "• Automated DNS Verification: Checks CNAME and A record configurations via Node DNS resolver.\n"
        "• Seamless Edge Rewriting: Middleware rewrites requests to the appropriate user magnet page without URL redirection or browser address bar changes.\n"
        "• Brand Customization: Custom logos, brand hex accents, dark/light theme overrides, and custom OpenGraph preview images."
    )

    # 4.10 Module 10: Real-Time Analytics & Telemetry
    h2 = doc.add_heading(level=2)
    h2.add_run("4.10 Module 10: Real-Time Analytics & Telemetry Engine")
    doc.add_paragraph(
        "• Core Conversion Metrics: Aggregates Total Views, Unique Visitors, Total Signups, and Conversion Rate %.\n"
        "• A/B Split Test Comparisons: Direct visual comparison of Variant A vs. Variant B conversion efficiency.\n"
        "• Device & Referrer Telemetry: Tracks desktop vs. mobile traffic distribution and top referral domains.\n"
        "• Granular Email Telemetry: Tracks `sent`, `delivered`, `opened`, `clicked`, `bounced`, and `complained` events using `EmailEvent` schema."
    )

    # =========================================================================
    # SECTION 5: Database Schema & Data Models
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("5. Database Schema & Data Models")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    doc.add_paragraph(
        "The application utilizes MongoDB with Mongoose ODM. Below are the key data entities and their structural specifications:"
    )

    schema_overview = [
        ["Model Name", "Primary Purpose", "Key Fields & Indexes"],
        ["Account", "User profiles, brand settings, credentials, and LinkedIn campaigns", "email (unique, index), username (unique), plan, brandColor, customDomain, domainVerified, linkedinPostCampaigns"],
        ["MagnetPage", "Landing page configs, copy, templates, A/B variants, PDF gates", "id, userEmail (index), slug (index), status, views, signups, template, pdfPages, pdfFreePages, afterSignupOption"],
        ["Lead", "Captured subscriber records, funnel attribution, and status", "id, userEmail (index), pageId (index), email (index), status, source, deviceType, referrer, signedUpAt (index)"],
        ["Sequence", "Email nurture drip campaigns linked to magnet pages", "id, userEmail (index), pageId, status, emails (array of steps with delays), stopOnBooking, stats"],
        ["Integration", "Third-party connector configurations and status", "id, userEmail (index), name, category, connected, trigger"],
        ["Resource", "Hosted digital assets, files, and deliverable links", "id (unique), userEmail (index), name, size, url, uploadedAt, fileExt"],
        ["EmailEvent", "Real-time granular email telemetry and webhook logs", "id (unique), userEmail (index), leadId, sequenceId, eventType (index), createdAt (index)"],
        ["PdfOtp", "Expiring 6-digit OTP verification codes for locked PDFs", "email, magnetId, code, token (unique), expiresAt (TTL index: 0s), used"]
    ]

    create_styled_table(doc, schema_overview[0], schema_overview[1:], col_widths=[1.5, 2.2, 2.8])

    # =========================================================================
    # SECTION 6: API Contracts & Webhook Architecture
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("6. API Contracts & Webhook Architecture")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    api_endpoints = [
        ["Method & Route", "Description", "Auth / Protection", "Payload / Response"],
        ["POST /api/auth/register", "Creates a new user account with hashed password", "Rate-limited (Upstash)", "Req: {name, email, password} -> Res: {success, user}"],
        ["POST /api/auth/login", "Authenticates user and issues HMAC session cookie", "Rate-limited (Upstash)", "Req: {email, password} -> Res: {success, token}"],
        ["POST /api/ai/generate-magnet", "Generates full magnet copy using Gemini AI", "JWT Session", "Req: {topic, targetAudience} -> Res: {headline, pitch, bullets, emailBody}"],
        ["POST /api/ai/optimize", "Optimizes existing copy hooks for higher conversion", "JWT Session", "Req: {text, tone} -> Res: {optimizedText}"],
        ["POST /api/pdf-gate/send-otp", "Dispatches 6-digit verification code for locked PDF", "Public Rate-limit", "Req: {email, magnetId, name} -> Res: {success, token}"],
        ["POST /api/pdf-gate/verify-otp", "Validates OTP code and unlocks full PDF reader", "Public", "Req: {token, code} -> Res: {verified: true, leadId}"],
        ["POST /api/upload", "Uploads assets and images to Cloudinary / Vercel Blob", "JWT Session", "Multipart Form Data -> Res: {url, publicId, size}"],
        ["POST /api/domain/verify", "Checks CNAME / A-records for custom domain", "JWT Session", "Req: {domain} -> Res: {verified, cnameValid, sslStatus}"],
        ["POST /api/track", "Records page views and conversion events", "Public", "Req: {magnetId, eventType, isVariantB, referrer} -> Res: {ok}"],
        ["POST /api/webhooks/booking", "Calendly webhook to halt active email sequences", "Webhook Signature", "Req: Calendly Invitee Event -> Halts sequence for lead"],
        ["POST /api/webhooks/email", "Resend webhook tracking open/click/bounce events", "Resend Secret", "Req: Resend Email Event -> Updates EmailEvent collection"],
        ["GET /api/cron/process-sequences", "Cron runner to process scheduled email sequence steps", "CRON_SECRET", "Executes due sequence emails across all accounts"],
        ["GET /api/cron/linkedin-sync", "Cron runner to poll LinkedIn comments and send DMs", "CRON_SECRET", "Dispatches DMs for matching trigger keywords"]
    ]

    create_styled_table(doc, api_endpoints[0], api_endpoints[1:], col_widths=[2.1, 2.0, 1.2, 1.2])

    # =========================================================================
    # SECTION 7: UI/UX Design System & Interactive Visuals
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("7. UI/UX Design System & Interactive Visuals")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    doc.add_paragraph(
        "LeadMagnets delivers a modern, high-polish user interface designed to maximize user engagement and delight:\n"
        "• Color Architecture: Curated palette featuring a deep brand cobalt blue (`#0066B2`), neutral zinc grays (`#09090b` to `#f4f4f5`), and crisp emerald accents for conversion states.\n"
        "• Dark / Light Theme Engine: Full multi-theme support with instant CSS token switching and highlight intensity controls.\n"
        "• 3D Interactive Canvas: Three.js and OGL-powered interactive 3D magnet canvas with floating geometric glass prisms and responsive aurora lighting.\n"
        "• Motion & Physics: Framer Motion layout transitions, spring physics on modal entries, and Lenis smooth scrolling for luxury editorial feel.\n"
        "• Typography: Google Fonts Geist and Outfit sans-serif hierarchy for ultra-clean readability across high-DPI desktop and mobile screens."
    )

    # =========================================================================
    # SECTION 8: Non-Functional Requirements (NFRs)
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("8. Non-Functional Requirements (NFRs)")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    nfr_data = [
        ["Category", "Standard / Target Benchmark", "Architectural Implementation"],
        ["Page Load Latency", "LCP < 1.2s, FID < 50ms, CLS = 0", "Edge SSR, automatic image optimization via Next/Image, critter inlined CSS, code splitting."],
        ["Edge Session Auth", "< 5ms token verification overhead", "Web Crypto API (crypto.subtle) HMAC-SHA256 executing directly inside Edge Middleware."],
        ["Security & Encryption", "AES-256 at rest, TLS 1.3 in transit", "Bcrypt password hashing, environment variable isolation, parameterized MongoDB queries."],
        ["Abuse & Bot Prevention", "Upstash Token Bucket Rate Limiting", "Strict rate limits on auth, password resets, OTP requests, and AI generation endpoints."],
        ["Data Privacy & Compliance", "GDPR, CCPA, and CAN-SPAM compliant", "One-click unsubscribe headers, explicit consent checkboxes, and data purge capabilities."],
        ["Availability & Uptime", "99.9% Target Service Level", "Stateless serverless Next.js functions on Vercel paired with MongoDB Atlas auto-scaling clusters."]
    ]

    create_styled_table(doc, nfr_data[0], nfr_data[1:], col_widths=[1.8, 2.3, 2.4])

    # =========================================================================
    # SECTION 9: Quality Assurance & Acceptance Criteria
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("9. Quality Assurance & Acceptance Criteria")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    test_scenarios = [
        ["Test ID", "Feature / Flow", "Test Condition", "Expected Outcome"],
        ["TC-01", "Edge Session Validation", "Access `/dashboard` without session token cookie", "Instant 307 redirect to `/login?from=/dashboard` without flashing dashboard UI."],
        ["TC-02", "Lead Magnet Creation", "Create magnet with custom prompt, bullets, and deliverable", "Magnet saves to MongoDB; public slug renders identical layout in < 1s."],
        ["TC-03", "Locked PDF OTP Gate", "Enter email on page 3 of locked PDF", "6-digit OTP received in inbox; verifying unlocks full PDF reader in-place without page reload."],
        ["TC-04", "Sequence Booking Stop", "Lead signs up, receives Email #1, then books Calendly call", "Calendly webhook updates Lead status to `stopped`; Email #2 is canceled."],
        ["TC-05", "LinkedIn Comment Trigger", "Comment 'resource' on configured LinkedIn campaign post", "System detects comment via cron/webhook and sends automated DM with magnet URL in < 60s."],
        ["TC-06", "A/B Split Test Telemetry", "100 unique visitors view magnet page with Variant B enabled", "Traffic is split 50/50; conversion counters accurately record Variant A vs B rates."],
        ["TC-07", "Custom Domain Routing", "CNAME verified domain `get.mybrand.com` requests root path", "Middleware rewrites request to user's designated magnet page seamlessly."]
    ]

    create_styled_table(doc, test_scenarios[0], test_scenarios[1:], col_widths=[0.8, 1.8, 2.0, 1.9])

    # =========================================================================
    # SECTION 10: Product Roadmap & Release Strategy
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("10. Product Roadmap & Future Horizons")
    r.font.name = "Arial"
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 102, 178)

    roadmap_data = [
        ["Milestone / Phase", "Target Scope & Core Features", "Delivery Horizon"],
        ["Phase 1: Core Foundation (Current)", "Full Lead Magnet builder, Locked PDF OTP gate, NextAuth+HMAC auth, MongoDB persistence, Resend email delivery, Tiptap editor.", "Completed & Production-Ready"],
        ["Phase 2: Growth & Social Engine (Current)", "LinkedIn comment automation, Calendly stop-on-booking webhooks, Google Gemini AI generator, A/B split testing, Pipedrive CRM sync.", "Active Deployment"],
        ["Phase 3: Multi-Channel Expansion", "Twitter / X Auto-DM integration, Instagram comment triggers, WhatsApp business deliverable delivery, Notion database auto-sync.", "Q4 2026"],
        ["Phase 4: Enterprise & Team Workspaces", "Multi-user role-based access control (RBAC), agency client sub-accounts, audit logging, advanced funnel cohort retention analytics.", "Q1 2027"]
    ]

    create_styled_table(doc, roadmap_data[0], roadmap_data[1:], col_widths=[1.8, 3.2, 1.5])

    # Output File
    output_path = os.path.join(os.getcwd(), "LeadMagnets_Project_PRD.docx")
    doc.save(output_path)
    print(f"Successfully generated PRD document at: {output_path}")

if __name__ == "__main__":
    build_complete_prd()
